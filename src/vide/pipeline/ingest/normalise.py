"""Turn whatever the client sent into plain text.

Scripts arrive as .docx from writers, .pdf from agents, .html from Google Docs
exports and .xlsx from whoever was tracking episodes in a spreadsheet. Each
loader returns text plus whatever structure it could preserve; everything
downstream works on the text.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import structlog

log = structlog.get_logger(__name__)

#: Characters that arrive from word processors and break naive matching.
_SUBSTITUTIONS = {
    "‘": "'", "’": "'",      # curly single quotes
    "“": '"', "”": '"',      # curly double quotes
    " ": " ",                      # non-breaking space
    "​": "", "﻿": "",        # zero-width space, BOM
    "…": "...",                    # ellipsis
}


class UnsupportedFormat(ValueError):
    pass


@dataclass(slots=True)
class NormalisedDocument:
    text: str
    source_filename: str
    content_type: str
    #: Anything the loader noticed but could not resolve — surfaced, not guessed.
    notes: list[str] = field(default_factory=list)

    @property
    def lines(self) -> list[str]:
        return self.text.splitlines()


def clean_text(raw: str) -> str:
    """Normalise whitespace and typographic characters.

    Dashes are deliberately left alone: the en dash is the field separator in
    this house's scene headers, so collapsing it to a hyphen would destroy the
    structure the parser depends on.
    """
    for bad, good in _SUBSTITUTIONS.items():
        raw = raw.replace(bad, good)
    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    raw = "\n".join(line.rstrip() for line in raw.split("\n"))
    return re.sub(r"\n{4,}", "\n\n\n", raw).strip() + "\n"


def _load_text(data: bytes) -> tuple[str, list[str]]:
    return data.decode("utf-8", errors="replace"), []


def _load_docx(data: bytes) -> tuple[str, list[str]]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    notes = []
    if document.tables:
        notes.append(
            f"{len(document.tables)} table(s) present; rows appended after the prose"
        )
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
    return "\n".join(parts), notes


def _load_pdf(data: bytes) -> tuple[str, list[str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    notes = []
    empty = sum(1 for p in pages if not p.strip())
    if empty:
        notes.append(
            f"{empty} of {len(pages)} pages yielded no text — likely scanned; "
            "OCR would be needed to read them"
        )
    return "\n\n".join(pages), notes


def _load_html(data: bytes) -> tuple[str, list[str]]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    notes = []
    tables = soup.find_all("table")
    if tables:
        # A Google Sheets export is a table, not prose — keep the grid.
        rows = []
        for table in tables:
            for tr in table.find_all("tr"):
                cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
                if any(cells):
                    rows.append(" | ".join(cells))
        if rows:
            notes.append(f"parsed {len(tables)} table(s) as pipe-delimited rows")
            return "\n".join(rows), notes

    return soup.get_text("\n"), notes


def _load_xlsx(data: bytes) -> tuple[str, list[str]]:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts, notes = [], []
    for sheet in book.worksheets:
        parts.append(f"## SHEET: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    notes.append(f"{len(book.worksheets)} sheet(s): {', '.join(s.title for s in book.worksheets)}")
    return "\n".join(parts), notes


_LOADERS = {
    ".md": _load_text,
    ".txt": _load_text,
    ".text": _load_text,
    ".fountain": _load_text,
    ".docx": _load_docx,
    ".pdf": _load_pdf,
    ".html": _load_html,
    ".htm": _load_html,
    ".xlsx": _load_xlsx,
    ".xlsm": _load_xlsx,
}

_CONTENT_TYPES = {
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".html": "text/html",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def normalise(path: Path | str) -> NormalisedDocument:
    path = Path(path)
    return normalise_bytes(path.read_bytes(), path.name)


def normalise_bytes(data: bytes, filename: str) -> NormalisedDocument:
    suffix = Path(filename).suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise UnsupportedFormat(
            f"{filename}: no loader for {suffix!r}; supported: "
            f"{', '.join(sorted(_LOADERS))}"
        )

    raw, notes = loader(data)
    text = clean_text(raw)
    if not text.strip():
        notes.append("document produced no text")

    log.info(
        "ingest.normalised",
        filename=filename,
        chars=len(text),
        lines=text.count("\n"),
        notes=len(notes),
    )
    return NormalisedDocument(
        text=text,
        source_filename=filename,
        content_type=_CONTENT_TYPES.get(suffix, "application/octet-stream"),
        notes=notes,
    )
